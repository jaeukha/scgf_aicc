"""Mock 프로필 end-to-end 점검. 실행: python -m pytest -q"""
import io
import json
import os
import tempfile
import wave
from pathlib import Path

DB = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ.setdefault("DATABASE_URL", f"sqlite:///{DB}")  # 빈 Postgres DB를 지정하면 Postgres에서 실행
os.environ.update(MOCK_DELAY_SECONDS="0", COACH_MODE="mock", PROVIDER_MODE="mock")

import pytest  # noqa: E402
from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

command.upgrade(Config("alembic.ini"), "head")  # 마이그레이션 자체도 검증

from aicc import main, providers  # noqa: E402
from aicc.db import ApiCallTrace, Session  # noqa: E402


@pytest.fixture(scope="module")
def c():
    with TestClient(main.app) as client:
        yield client


def new_session(c, **kw):
    return c.post("/api/v1/sessions", json=kw).json()["session_id"]


def say(c, sid, text, speaker="customer", **kw):
    r = c.post(f"/api/v1/sessions/{sid}/simulate", json={"text": text, "speaker": speaker, **kw})
    assert r.status_code == 200, r.text
    return r.json()


def test_core_flow_trace_and_crm(c):
    emp = c.get("/api/v1/crm/employees").json()[0]["employee_id"]
    cust = c.get("/api/v1/crm/customers").json()["items"][0]["customer_id"]
    sid = new_session(c)
    for line in providers.MOCK_SCRIPTS["extension"]:
        out = say(c, sid, line["text"], line["speaker"])
    assert out["coach_status"] == "success"
    assert out["routing"]["paths"][0]["topic"] == "만기·연장"
    s = c.get(f"/api/v1/sessions/{sid}").json()
    assert s["current_topic"] == "만기·연장"

    types = {e["type"] for e in c.get(f"/api/v1/sessions/{sid}/events").json()}
    assert {"transcript", "analysis.routing", "analysis.sentiment_text", "coach"} <= types

    # trace: 한 요청에 backend → db/decider/knowledge/coach span이 부모-자식으로 묶인다
    r = c.post(f"/api/v1/sessions/{sid}/simulate", json={"text": "감사합니다", "speaker": "customer"})
    spans = c.get(f"/api/v1/monitoring/traces/{r.headers['x-trace-id']}").json()
    ops = {s["operation"] for s in spans}
    assert {"decider.routing", "knowledge.search", "coach.generate"} <= ops
    root = next(s for s in spans if s["service_name"] == "backend")
    assert root["call_kind"] == "http" and all(s["parent_span_id"] == root["span_id"] for s in spans if s is not root)

    # Coach 장애 주입: 전사·Decider 결과는 유지되고 실패가 기록된다
    out = say(c, sid, "만기가 지났어요", inject_error={"coach": "timeout"})
    assert out["coach_status"] == "failed" and out["routing"]["paths"]
    errs = c.get("/api/v1/monitoring/errors").json()
    assert any(e["error_code"] == "UPSTREAM_TIMEOUT" and e["service_name"] == "coach" for e in errs)
    ov = c.get("/api/v1/monitoring/overview").json()
    assert ov["by_service"]["coach"]["timeouts"] >= 1 and ov["last_5m"]["p95_ms"] is not None
    series = c.get("/api/v1/monitoring/series?minutes=5").json()
    assert len(series["buckets"]) == 60 and series["summary"]["requests"] >= 1
    assert sum(x["requests"] for x in series["buckets"]) == series["summary"]["requests"]

    # 종료 → AI 요약, 직원 검토는 AI 원본을 덮어쓰지 않음
    ended = c.post(f"/api/v1/sessions/{sid}/end").json()
    assert ended["status"] == "completed" and ended["ai_summary"].startswith("[AI 생성 요약")
    assert all(section in ended["ai_summary"] for section in (
        "## 상담 목적", "## 대화에서 확인된 사실", "## 상담 경과", "## 미확인 항목", "## 후속 업무"))
    assert c.post(f"/api/v1/sessions/{sid}/simulate", json={"text": "x"}).json()["error"]["code"] == "SESSION_NOT_ACTIVE"
    h = {"X-Employee-Id": emp}
    rv = c.patch(f"/api/v1/crm/sessions/{sid}/review", headers=h,
                 json={"review_status": "corrected", "final_summary": "직원 확정 요약", "final_topic": "만기·연장"}).json()
    assert rv["ai_summary"] == ended["ai_summary"] and rv["final_summary"] == "직원 확정 요약"
    assert c.patch(f"/api/v1/crm/sessions/{sid}/review", json={"review_status": "x"}).status_code == 422
    c.patch(f"/api/v1/crm/sessions/{sid}/assignment", headers=h, json={"assigned_employee_id": emp, "customer_id": cust})
    assert c.post(f"/api/v1/crm/sessions/{sid}/notes", json={"note_text": "메모", "author_employee_id": emp}).status_code == 201

    # AI 제안 → 상담원 선택 시에만 후속 업무 생성
    action = next(e for e in c.get(f"/api/v1/sessions/{sid}/events").json()
                  if e["type"] == "coach" and e["data"]["action_items"])["data"]["action_items"][0]
    t = c.post("/api/v1/crm/tasks", headers=h, json={"session_id": sid, "title": action, "source": "ai_suggested",
                                                      "assigned_employee_id": emp, "due_at": "2000-01-01T00:00:00Z"}).json()
    assert t["overdue"] and t["customer_id"] == cust
    assert c.get("/api/v1/crm/tasks?overdue_only=true").json()
    t = c.patch(f"/api/v1/crm/tasks/{t['task_id']}", headers=h, json={"status": "completed"}).json()
    assert t["completed_at"] and not t["overdue"]
    assert c.post("/api/v1/crm/tasks", json={"session_id": sid, "title": "x", "bogus": 1}).status_code == 422

    detail = c.get(f"/api/v1/crm/sessions/{sid}").json()
    assert {a["action"] for a in detail["audit"]} == {"review", "assign", "note_add"}
    assert c.get(f"/api/v1/crm/customers/{cust}").json()["timeline"][0]["session_id"] == sid
    assert c.get(f"/api/v1/crm/sessions?employee_id={emp}&task_status=completed").json()["total"] == 1

    # trace에는 전사 원문이 들어가지 않는다
    with Session() as db:
        dump = " ".join(str(v) for r in db.query(ApiCallTrace) for v in r.__dict__.values())
    assert "만기가 다가와서" not in dump and "감사합니다" not in dump


def test_documents_routes_two_topics(c):
    sid = new_session(c)
    for line in providers.MOCK_SCRIPTS["documents"]:
        out = say(c, sid, line["text"], line["speaker"])
    assert {p["topic"] for p in out["routing"]["paths"]} == {"보증 신청", "서류 문의"}


def test_delay_callback_route_does_not_ask_for_fund_timing():
    progress = providers.ROUTING_CONFIG["nodes"]["progress_kind"]["branches"]["지연·재연락 요청"]
    assert progress["next"] == "callback"
    assert "접수일" in providers.ROUTING_CONFIG["nodes"]["callback"]["action"]
    assert "prog_urgency" not in providers.ROUTING_CONFIG["nodes"]


def test_closing_summary_eval_dataset(c):
    """검증 세트는 자유서술 요약 모델을 도입할 때도 같은 사실·금지문구 기준으로 재사용한다."""
    dataset = json.loads((Path(__file__).parents[1] / "data" / "evals" / "closing_summary_v1.json").read_text(encoding="utf-8"))
    for case in dataset["cases"]:
        sid = new_session(c)
        route = {"paths": []}
        for line in case["transcript"]:
            route = say(c, sid, line["text"], line["speaker"])["routing"]
        if case["register_route_action"]:
            action = next(p["action"] for p in route["paths"] if p.get("action"))
            created = c.post("/api/v1/crm/tasks", json={"session_id": sid, "title": action, "source": "ai_suggested"})
            assert created.status_code == 201, case["id"]
        summary = c.post(f"/api/v1/sessions/{sid}/end").json()["ai_summary"]
        for text in case["must_include"]:
            assert text in summary, f"{case['id']}: missing {text!r}\n{summary}"
        for text in case["must_not_include"]:
            assert text not in summary, f"{case['id']}: invented {text!r}\n{summary}"
        if case["register_route_action"]:
            assert action in summary, f"{case['id']}: registered action missing\n{summary}"


def test_stale_coach_discarded(c, monkeypatch):
    sid = new_session(c)
    real = providers.generate_guidance

    def slow(context, routing):  # 생성 도중 새 발화가 들어온 상황
        main._generation[sid] += 1
        return real(context, routing)
    monkeypatch.setattr(providers, "generate_guidance", slow)
    assert say(c, sid, "만기 연장하고 싶어요")["coach_status"] == "stale"


def test_audio_upload_mock_asr(c):
    sid = new_session(c, input_mode="file")
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1), w.setsampwidth(2), w.setframerate(16000), w.writeframes(b"\0" * 32000)
    r = c.post(f"/api/v1/sessions/{sid}/audio?scenario=progress", files={"file": ("a.wav", buf.getvalue())})
    assert r.json()["audio_seconds"] == 1.0
    assert providers.validate_audio("recording.m4a", b"m4a") is None
    import time
    for _ in range(50):  # 파일 재생은 전용 스레드에서 진행
        finals = [e for e in c.get(f"/api/v1/sessions/{sid}/events").json() if e["type"] == "transcript"]
        if len(finals) == 3:
            break
        time.sleep(0.1)
    assert len(finals) == 3 and finals[0]["data"]["source"] == "asr"
    bad = c.post(f"/api/v1/sessions/{sid}/audio", files={"file": ("a.txt", b"hi")})
    assert bad.status_code == 415 and bad.json()["error"]["code"] == "AUDIO_FORMAT_UNSUPPORTED" and bad.json()["trace_id"]
    assert c.get("/api/v1/sessions/00000000-0000-0000-0000-000000000000").json()["error"]["code"] == "SESSION_NOT_FOUND"


def test_websocket_stream(c):
    sid = new_session(c)
    say(c, sid, "서류 발급은 어디서 하나요?")
    with c.websocket_connect(f"/api/v1/sessions/{sid}/stream") as ws:
        got = [ws.receive_json()["type"] for _ in range(6)]
    assert "transcript" in got and "coach" in got


def test_health(c):
    assert c.get("/health/live").json()["status"] == "ok"
    assert c.get("/health/ready").json()["db"] is True
    assert c.get("/").status_code == 200


def test_http_decider_adapter(monkeypatch):
    """PROVIDER_MODE=local 경로: 연결 실패 1회 재시도, trace 헤더 전파, http span 기록, 응답 형식 그대로 라우팅."""
    import httpx
    from aicc import telemetry
    calls = []

    def fake_post(url, timeout, headers, json):
        calls.append(headers["X-Trace-Id"])
        if len(calls) == 1:
            raise httpx.ConnectError("down")
        result, _ = providers.MockDecider().decide(json["state"], json["questions"], [], json["layout"], None)
        return httpx.Response(200, json=result | {"seconds": 0.01}, request=httpx.Request("POST", url))
    monkeypatch.setattr(providers.httpx, "post", fake_post)
    tid = telemetry.new_trace()
    r = providers.legacy.run_routing(providers.HttpDecider(), "[고객] 만기가 다가와서 연장하고 싶어요", providers.ROUTING_CONFIG)
    assert r["paths"][0]["topic"] == "만기·연장" and set(calls) == {tid}
    with Session() as db:
        spans = db.query(ApiCallTrace).filter_by(trace_id=tid).all()
    assert all(s.call_kind == "http" and s.http_status_code == 200 for s in spans)
    assert spans[0].retry_count == 1


def test_mic_segment_mock(c):
    sid = new_session(c, input_mode="microphone")
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1), w.setsampwidth(2), w.setframerate(48000), w.writeframes(b"\0" * 96000)
    out = c.post(f"/api/v1/sessions/{sid}/audio-segment?speaker=customer", files={"file": ("seg.wav", buf.getvalue())}).json()
    assert len(out) == 1
    ev = [e for e in c.get(f"/api/v1/sessions/{sid}/events").json() if e["type"] == "transcript"][0]
    assert ev["data"]["speaker"] == "customer" and "Mock ASR" in ev["data"]["text"]


def test_twenty_sync_idempotent(c, monkeypatch):
    """AICC→Twenty 단방향: 첫 동기화는 POST, 재동기화는 PATCH, Twenty에서 삭제(404)되면 재생성, 실패는 last_error."""
    import httpx
    from aicc import config, twenty_sync
    from aicc.db import CrmSyncLink
    monkeypatch.setattr(config, "TWENTY_ENABLED", True)
    import itertools
    store, calls, ids = {}, [], itertools.count()

    def fake(method, url, json, headers, timeout):
        path = url.split("/rest/")[1]
        calls.append((method, path.split("/")[0]))
        req = httpx.Request(method, url)
        if method == "POST":
            new = f"00000000-0000-0000-0000-{next(ids):012d}"
            store[new] = json
            return httpx.Response(201, json={"data": {"create": {"id": new}}}, request=req)
        tid = path.split("/")[1]
        return httpx.Response(200 if tid in store else 404, json={}, request=req)
    monkeypatch.setattr(twenty_sync.httpx, "request", fake)
    monkeypatch.setattr(twenty_sync, "schedule", lambda sid: None)

    cust = c.get("/api/v1/crm/customers").json()["items"][1]["customer_id"]
    sid = new_session(c, customer_id=cust)
    c.post("/api/v1/crm/tasks", json={"session_id": sid, "title": "만기일 확인"})
    first = c.post(f"/api/v1/crm/sessions/{sid}/sync-twenty").json()
    assert [m for m, _ in calls] == ["POST", "POST", "POST", "POST"]  # company, consultation, task, taskTarget
    calls.clear()
    c.post(f"/api/v1/crm/sessions/{sid}/sync-twenty")
    assert {m for m, _ in calls} == {"PATCH"}  # 중복 생성 없음
    del store[first["consultation_id"]]  # Twenty에서 삭제됨
    assert c.post(f"/api/v1/crm/sessions/{sid}/sync-twenty").json()["consultation_id"] != first["consultation_id"]

    # 동시 동기화(종료·검토·업무 변경 직후)에도 업무는 한 번만 생성, 마감은 UTC로 변환
    import threading
    sid2 = new_session(c, customer_id=cust)
    c.post("/api/v1/crm/tasks", json={"session_id": sid2, "title": "동시", "due_at": "2026-10-31T09:00:00+09:00"})
    calls.clear()
    th = [threading.Thread(target=twenty_sync.sync_session, args=(main.uuid.UUID(sid2),)) for _ in range(4)]
    [t.start() for t in th], [t.join() for t in th]
    assert calls.count(("POST", "tasks")) == 1
    assert [v.get("dueAt") for v in store.values() if v.get("title") == "동시"] == ["2026-10-31T00:00:00+00:00"], store

    def down(*a, **k):
        raise httpx.ConnectError("down")
    monkeypatch.setattr(twenty_sync.httpx, "request", down)
    assert c.post(f"/api/v1/crm/sessions/{sid}/sync-twenty").json()["error"]["code"] == "UPSTREAM_UNAVAILABLE"
    with Session() as db:
        assert db.query(CrmSyncLink).filter_by(entity_type="customer", aicc_id=cust).one().last_error


def test_pii_masked_before_storage(c):
    sid = new_session(c)
    say(c, sid, "사업자번호는 123-45-67890이고 전화는 010-1234-5678, 신청일은 2026-10-10이에요")
    text = next(e for e in c.get(f"/api/v1/sessions/{sid}/events").json() if e["type"] == "transcript")["data"]["text"]
    assert "[사업자등록번호]" in text and "[전화번호]" in text and "2026-10-10" in text and "67890" not in text


def test_robustness_errors(c):
    import uuid as u
    bad = str(u.uuid4())
    assert c.post("/api/v1/sessions", json={"assigned_employee_id": bad}).json()["error"]["code"] == "INVALID_REFERENCE"
    assert c.post("/api/v1/sessions", json={"zzz": 1}).status_code == 422
    sid = new_session(c)
    assert c.post(f"/api/v1/sessions/{sid}/simulate", json={"text": "   "}).status_code == 422
    t = c.post("/api/v1/crm/tasks", json={"session_id": sid, "title": "x"}).json()
    r = c.patch(f"/api/v1/crm/tasks/{t['task_id']}", json={"status": None})
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_REFERENCE"
    assert c.get("/api/v1/crm/customers?size=-1").status_code == 422
    assert c.get("/api/v1/nope").json()["error"]["code"] == "NOT_FOUND"
    import pytest as pt
    from starlette.websockets import WebSocketDisconnect
    with pt.raises(WebSocketDisconnect):
        with c.websocket_connect(f"/api/v1/sessions/{sid}/stream", headers={"origin": "http://evil.example"}) as ws:
            ws.receive_json()


def test_rag_ingest_and_search(tmp_path, monkeypatch):
    """사내 지식베이스: 파싱·조각·색인(멱등)·변경 반영·삭제 반영·검색·상담 흐름 연결. mock 임베딩 + 메모리 Qdrant."""
    import shutil
    import docx
    from aicc import config, knowledge
    monkeypatch.setattr(config, "QDRANT_URL", ":memory:")
    monkeypatch.setattr(config, "EMBED_BASE_URL", "mock")
    monkeypatch.setattr(config, "QDRANT_COLLECTION", "test_kb")
    monkeypatch.setattr(config, "KNOWLEDGE_MIN_SCORE", 0.1)
    monkeypatch.setattr(knowledge, "_client", None)
    for f in (config.ROOT / "data" / "kb_samples").iterdir():
        shutil.copy(f, tmp_path)
    d = docx.Document()
    d.add_heading("서류 보완", level=1)
    d.add_paragraph("보완 요청을 받은 고객에게는 업무시스템에서 요청 내용과 제출 기한을 확인해 안내한다.")
    d.save(tmp_path / "서류보완.docx")
    (tmp_path / "옛문서.hwp").write_bytes(b"x")

    first = {r["document"]: r for r in knowledge.ingest_dir(tmp_path)}
    assert first["서류보완.docx"]["status"] == "indexed" and first["FAQ_보증료.csv"]["chunks"] == 2
    assert first["옛문서.hwp"]["status"] == "failed"
    assert all(r["status"] == "unchanged" for r in knowledge.ingest_dir(tmp_path) if r["document"] != "옛문서.hwp")

    hits = knowledge.search("[고객] 이번 달 연체가 됐는데 분할상환 되나요", ["상환·연체"])
    assert hits and hits[0]["title"] == "연체_분할상환_안내" and hits[0]["section"] == "연체 고객 응대"

    (tmp_path / "서류보완.docx").unlink()
    gone = {r["document"]: r["status"] for r in knowledge.ingest_dir(tmp_path)}
    assert gone["서류보완.docx"] == "removed"

    # 상담 흐름에서 RAG 결과가 참고 자료(coach 근거)로 들어간다
    monkeypatch.setattr(config, "KNOWLEDGE_MODE", "qdrant")
    with TestClient(main.app) as c:
        sid = new_session(c)
        say(c, sid, "보증 받은 대출이 연체가 됐어요. 나눠서 갚고 싶어요")
        refs = next(e for e in c.get(f"/api/v1/sessions/{sid}/events").json() if e["type"] == "analysis.knowledge")["data"]["result"]["references"]
    assert refs and refs[0]["source"].endswith(".md") and "score" in refs[0]
