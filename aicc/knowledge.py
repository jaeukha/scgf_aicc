"""사내 지식베이스 RAG: 문서 파싱 → 조각 나누기 → EmbeddingGemma 2 임베딩 → Qdrant 색인/검색.

색인:  python -m aicc.knowledge ingest data/kb_docs      (바뀐 문서만 다시 색인, 지워진 문서는 제거)
확인:  python -m aicc.knowledge stats | search "만기 연장 서류"
검색 결과는 근거 표시용 메타데이터(document_id, title, section, page, source)를 함께 돌려준다.
문서 본문은 참고자료일 뿐 지시가 아니다(프롬프트 인젝션 방지는 Coach 프롬프트 쪽 책임).
"""
import hashlib
import re
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

import httpx

from . import config
from .telemetry import span, trace_var

SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".hwpx", ".csv", ".xlsx"}
CHUNK_CHARS, OVERLAP_CHARS = 700, 120  # 한국어 기준 약 300~400토큰. 문서가 오면 검색 평가로 조정


# ---------- 파싱: (본문, 섹션 제목, 페이지) 블록 목록 ----------
def _md_blocks(text):
    section = None
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)", para)
        if heading:
            section = heading.group(2).strip()
            para = para[heading.end():].strip()
            if not para:
                continue
        yield para, section, None


def _pdf_blocks(path):
    from pypdf import PdfReader
    for no, page in enumerate(PdfReader(str(path)).pages, 1):
        for para in re.split(r"\n\s*\n", page.extract_text() or ""):
            if para.strip():
                yield para.strip(), None, no


def _docx_blocks(path):
    import docx
    section = None
    for p in docx.Document(str(path)).paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style is not None and p.style.name.lower().startswith(("heading", "제목")):
            section = text
            continue
        yield text, section, None


def _hwpx_blocks(path):
    """HWPX(한글 2014+ 개방형 XML). 바이너리 .hwp는 HWPX/PDF로 변환해서 넣는다."""
    with zipfile.ZipFile(path) as z:
        for name in sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n)):
            root = ElementTree.fromstring(z.read(name))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t")).strip()
                    if text:
                        yield text, None, None


def _table_blocks(rows):
    """FAQ 표: 첫 행을 머리글로 보고 한 행을 한 블록("머리글: 값")으로."""
    rows = [[str(c).strip() if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return
    header = rows[0]
    for n, r in enumerate(rows[1:], 2):  # 행마다 따로 조각이 되도록 첫 칸(보통 질문)을 섹션으로
        yield "\n".join(f"{h}: {v}" for h, v in zip(header, r) if v), (r[0][:80] or f"{n}행"), None


def parse(path: Path):
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return list(_md_blocks(path.read_text(encoding="utf-8", errors="replace")))
    if ext == ".pdf":
        return list(_pdf_blocks(path))
    if ext == ".docx":
        return list(_docx_blocks(path))
    if ext == ".hwpx":
        return list(_hwpx_blocks(path))
    if ext == ".csv":
        import csv
        with open(path, encoding="utf-8-sig", newline="") as f:
            return list(_table_blocks(csv.reader(f)))
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return [b for ws in wb.worksheets for b in _table_blocks(ws.iter_rows(values_only=True))]
    raise ValueError(f"지원하지 않는 형식: {ext} (지원: {', '.join(sorted(SUPPORTED))})")


def chunk(blocks):
    """같은 섹션·페이지 안에서 블록을 이어 붙여 약 CHUNK_CHARS 크기로 자르고, 앞 조각 끝을 조금 겹친다."""
    out, buf, key = [], "", None
    def flush():
        if buf.strip():
            out.append({"text": buf.strip(), "section": key[0], "page": key[1]})
    for text, section, page in blocks:
        if key != (section, page) and buf:
            flush()
            buf = ""
        key = (section, page)
        for piece in [text[i:i + CHUNK_CHARS] for i in range(0, len(text), CHUNK_CHARS)]:  # 아주 긴 문단 분할
            if buf and len(buf) + len(piece) + 1 > CHUNK_CHARS:
                flush()
                buf = buf[-OVERLAP_CHARS:]
            buf = f"{buf}\n{piece}" if buf else piece
    if buf:
        flush()
    return out


# ---------- 임베딩 ----------
def _mock_vectors(texts):
    """테스트·모델 없는 환경용 결정적 임베딩(문자 2-gram 해시). 의미 검색 품질은 없다."""
    import math
    out = []
    for t in texts:
        v = [0.0] * config.EMBED_DIM
        t = re.sub(r"\s+", "", t.split("query: ")[-1].split("text: ")[-1])
        for i in range(len(t) - 1):
            v[int(hashlib.md5(t[i:i + 2].encode()).hexdigest(), 16) % config.EMBED_DIM] += 1
        n = math.sqrt(sum(x * x for x in v)) or 1
        out.append([x / n for x in v])
    return out


def embed(texts, kind, titles=None):
    if config.EMBED_BASE_URL == "mock":
        return _mock_vectors(texts)
    with span("knowledge", f"knowledge.embed_{kind}", provider="embeddinggemma-2", call_kind="http",
              model_version=config.EMBED_MODEL) as extra:
        r = httpx.post(config.EMBED_BASE_URL.rstrip("/") + "/embed", json={"texts": texts, "kind": kind, "titles": titles},
                       headers={"X-Trace-Id": trace_var.get() or ""}, timeout=config.EMBED_TIMEOUT_SECONDS)
        extra["http_status_code"] = r.status_code
        r.raise_for_status()
        body = r.json()
    if body["dim"] != config.EMBED_DIM:
        raise ValueError("임베딩 차원이 컬렉션 설정과 다릅니다.")
    return body["vectors"]


# ---------- Qdrant ----------
_client = None


def client():
    global _client
    if _client is None:
        from qdrant_client import QdrantClient
        _client = QdrantClient(location=":memory:") if config.QDRANT_URL == ":memory:" else \
            QdrantClient(url=config.QDRANT_URL, timeout=int(config.EMBED_TIMEOUT_SECONDS))
    return _client


def ensure_collection():
    from qdrant_client import models as m
    c = client()
    if not c.collection_exists(config.QDRANT_COLLECTION):
        c.create_collection(config.QDRANT_COLLECTION,
                            vectors_config=m.VectorParams(size=config.EMBED_DIM, distance=m.Distance.COSINE))
        if config.QDRANT_URL != ":memory:":  # 메모리 모드는 인덱스를 지원하지 않음
            for field in ("document_id", "doc_sha"):
                c.create_payload_index(config.QDRANT_COLLECTION, field, m.PayloadSchemaType.KEYWORD)
    return c


def _doc_filter(document_id):
    from qdrant_client import models as m
    return m.Filter(must=[m.FieldCondition(key="document_id", match=m.MatchValue(value=document_id))])


def ingest_file(path: Path, root: Path):
    """한 문서를 색인. 내용이 같으면 건너뛰고, 바뀌었으면 기존 조각을 지우고 다시 넣는다(멱등)."""
    from qdrant_client import models as m
    c = ensure_collection()
    rel = path.relative_to(root).as_posix()
    document_id = hashlib.sha1(rel.encode()).hexdigest()[:16]
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    same = c.count(config.QDRANT_COLLECTION, count_filter=m.Filter(must=[
        m.FieldCondition(key="document_id", match=m.MatchValue(value=document_id)),
        m.FieldCondition(key="doc_sha", match=m.MatchValue(value=sha))]), exact=True).count
    if same:
        return {"document": rel, "status": "unchanged", "chunks": same}
    pieces = chunk(parse(path))
    title = path.stem
    vectors = []
    for i in range(0, len(pieces), 32):
        batch = pieces[i:i + 32]
        vectors += embed([p["text"] for p in batch], "document", [title] * len(batch))
    c.delete(config.QDRANT_COLLECTION, points_selector=m.FilterSelector(filter=_doc_filter(document_id)))
    now = datetime.now(timezone.utc).isoformat()
    c.upsert(config.QDRANT_COLLECTION, points=[m.PointStruct(
        id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{document_id}:{n}")), vector=v,
        payload={"document_id": document_id, "doc_sha": sha, "title": title, "source": rel, "chunk_no": n,
                 "section": p["section"], "page": p["page"], "text": p["text"], "ingested_at": now})
        for n, (p, v) in enumerate(zip(pieces, vectors))])
    return {"document": rel, "status": "indexed", "chunks": len(pieces)}


def ingest_dir(root: Path):
    from qdrant_client import models as m
    root = root.resolve()
    files = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED)
    results = []
    for f in files:
        try:
            results.append(ingest_file(f, root))
        except Exception as exc:  # 한 문서 실패가 나머지 색인을 막지 않게
            results.append({"document": f.relative_to(root).as_posix(), "status": "failed", "error": str(exc)[:200]})
    # 폴더에서 사라진 문서는 색인에서도 제거
    keep = {hashlib.sha1(f.relative_to(root).as_posix().encode()).hexdigest()[:16] for f in files}
    c = ensure_collection()
    known, offset = set(), None
    while True:
        points, offset = c.scroll(config.QDRANT_COLLECTION, limit=256, offset=offset, with_payload=["document_id", "source"])
        known |= {(p.payload["document_id"], p.payload["source"]) for p in points}
        if offset is None:
            break
    for document_id, source in known:
        if document_id not in keep:
            c.delete(config.QDRANT_COLLECTION, points_selector=m.FilterSelector(filter=_doc_filter(document_id)))
            results.append({"document": source, "status": "removed"})
    skipped = [p.name for p in root.rglob("*") if p.is_file() and p.suffix.lower() == ".hwp"]
    results += [{"document": n, "status": "failed", "error": ".hwp는 지원하지 않습니다. HWPX 또는 PDF로 저장해 넣으세요."}
                for n in skipped]
    return results


def query_text(transcript, topics):
    """검색 질의: 최근 고객 발화 위주 + 파악된 주제. 전사 30줄 전체를 넣으면 질의가 흐려진다."""
    lines = [l for l in transcript.splitlines() if l.strip()]
    customer = [l for l in lines if l.startswith("[고객]")] or lines
    return " ".join(topics) + " " + " ".join(re.sub(r"^\[[^\]]+\]\s*", "", l) for l in customer[-3:])


def search(transcript, topics, limit=None):
    c = client()
    if not c.collection_exists(config.QDRANT_COLLECTION):
        return []
    vector = embed([query_text(transcript, topics)[:2000]], "query")[0]
    with span("knowledge", "knowledge.qdrant_search", provider="qdrant",
              call_kind="inproc" if config.QDRANT_URL == ":memory:" else "http"):
        hits = c.query_points(config.QDRANT_COLLECTION, query=vector, limit=limit or config.KNOWLEDGE_TOP_K,
                              score_threshold=config.KNOWLEDGE_MIN_SCORE, with_payload=True).points
    hits = [h for h in hits if h.score >= hits[0].score - config.KNOWLEDGE_MAX_GAP] if hits else []
    return [{"document_id": h.payload["document_id"], "title": h.payload["title"],
             "section": " · ".join(x for x in [h.payload.get("section"), f"p.{h.payload['page']}" if h.payload.get("page") else None] if x) or "본문",
             "snippet": h.payload["text"][:400], "source": h.payload["source"], "score": round(h.score, 3)} for h in hits]


def healthy():
    try:
        client().get_collections()
        if config.EMBED_BASE_URL != "mock":
            httpx.get(config.EMBED_BASE_URL.rstrip("/") + "/health", timeout=2).raise_for_status()
        return True
    except Exception:
        return False


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "ingest" and len(sys.argv) == 3:
        for r in ingest_dir(Path(sys.argv[2])):
            print(r)
    elif cmd == "stats":
        c = ensure_collection()
        print(config.QDRANT_COLLECTION, "조각 수:", c.count(config.QDRANT_COLLECTION, exact=True).count)
    elif cmd == "search" and len(sys.argv) >= 3:
        for r in search("[고객] " + " ".join(sys.argv[2:]), []):
            print(r["score"], r["title"], r["section"], "|", r["snippet"][:80].replace("\n", " "))
    else:
        print(__doc__)
